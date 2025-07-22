from typing import Any, Dict, Iterable, Tuple

import openpyxl

from .. import logger
from .ADocument import ADocument


class XlsDocument(ADocument):
    def iterate_raw_text(self) -> Iterable[Tuple[int, str, Dict[str, Any]]]:
        try:
            wb = openpyxl.load_workbook(self.get_abs_path(), read_only=True, data_only=True)
        except Exception:
            logger.warning("Error while reading the file. Skipping")
            return None, {"ocr_used": False}

        nb_sheets = len(wb.worksheets)
        logger.info(f"Reading {nb_sheets} pages excel file")
        avct = -1
        all_text = []
        for k_sheet, sheet in enumerate(wb.worksheets):
            new_avct = int(k_sheet / nb_sheets * 100 / 10)
            if new_avct != avct:
                logger.info(f"Lecture page {k_sheet+1}/{nb_sheets}")
                avct = new_avct

            for row in sheet.iter_rows(values_only=True):
                row_text = [str(cell) for cell in row if cell is not None]
                if row_text:
                    all_text.append(" ".join(row_text))
            yield k_sheet, "\n".join(all_text).strip(), {"ocr_used": False}
